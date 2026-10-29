#!/usr/bin/env python3
"""Régénère db/seed_data/nomenclature_induni.json depuis le classeur « Nomenclature INDUNI ».

Usage : python3 script/build_nomenclature_json.py <classeur.xlsx> [sortie.json]

Feuille « Base article » : seules les lignes dont la colonne A vaut « oui » (« Intégrer dans
l'eshop ») sont reprises. Colonnes : B n° compta, C-E segmentation 1-3, F n° INDUNI, G désignation
INDUNI, H fournisseur, I réf. fournisseur, J désignation fournisseur, K unité de prix 1,
L mots clefs pour la recherche (séparés par « ; »).
"""
import json, sys
from openpyxl import load_workbook

src = sys.argv[1]
dst = sys.argv[2] if len(sys.argv) > 2 else "db/seed_data/nomenclature_induni.json"
clean = lambda v: None if v is None or str(v).strip() == "" else str(v).strip()
out = []
for r in load_workbook(src, data_only=True)["Base article"].iter_rows(min_row=2, values_only=True):
    if str(r[0] or "").strip().lower() != "oui":
        continue
    out.append({
        "nCompta": clean(r[1]), "segmentation1": clean(r[2]), "segmentation2": clean(r[3]),
        "segmentation3": clean(r[4]), "induniReference": clean(r[5]), "induniDesignation": clean(r[6]),
        "supplierName": clean(r[7]), "supplierReference": clean(r[8]), "supplierDesignation": clean(r[9]),
        "priceUnit": clean(r[10]), "searchKeywords": clean(r[11]),
    })
with open(dst, "w", encoding="utf-8") as f:
    json.dump(out, f, ensure_ascii=False, indent=1)
    f.write("\n")
print(f"{len(out)} lignes -> {dst}")
